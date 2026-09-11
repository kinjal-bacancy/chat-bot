"""Builders for real documents to parse.

Generating genuine PDF and DOCX bytes rather than stubbing the parsers: the
whole risk in this stage is how real files behave, which a stub cannot show.
"""

import io

import docx as python_docx
from reportlab.pdfgen import canvas


def pdf_bytes(pages: list[list[str]]) -> bytes:
    """A PDF where each inner list is the lines drawn on one page."""
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=(612, 792))
    for lines in pages:
        y = 750
        for line in lines:
            pdf.drawString(72, y, line)
            y -= 20
        pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def docx_bytes(paragraphs: list[str], tables: list[list[list[str]]] | None = None) -> bytes:
    buffer = io.BytesIO()
    document = python_docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    for rows in tables or []:
        table = document.add_table(rows=len(rows), cols=len(rows[0]))
        for row_index, row in enumerate(rows):
            for cell_index, value in enumerate(row):
                table.cell(row_index, cell_index).text = value
    document.save(buffer)
    return buffer.getvalue()


def xlsx_bytes(sheets: dict[str, list[list]]) -> bytes:
    import openpyxl

    buffer = io.BytesIO()
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        worksheet = workbook.create_sheet(title=name)
        for row in rows:
            worksheet.append(row)
    workbook.save(buffer)
    return buffer.getvalue()

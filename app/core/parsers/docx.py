"""Word document extraction."""

from pathlib import Path

import docx

from app.core.parsers.base import NoExtractableText, Page, ParsedDocument, normalise


def _table_to_text(table) -> str:
    """Flatten a table to one tab-separated line per row.

    Tables carry real answers -- pricing, specs, dates -- so dropping them
    loses content a user will ask about. Tab separation keeps the row/column
    relationship legible to the model without inventing a markup format.
    """
    rows = []
    for row in table.rows:
        cells = [cell.text.strip() for cell in row.cells]
        if any(cells):
            rows.append("\t".join(cells))
    return "\n".join(rows)


def parse(path: Path) -> ParsedDocument:
    document = docx.Document(str(path))

    blocks = [p.text for p in document.paragraphs if p.text.strip()]
    blocks.extend(filter(None, (_table_to_text(t) for t in document.tables)))

    # .docx has no reliable page boundaries -- pagination is decided by the
    # renderer, not stored in the file -- so the whole document is one page.
    text = normalise("\n\n".join(blocks))

    if not text:
        raise NoExtractableText("The document contains no text.")

    return ParsedDocument(pages=[Page(number=1, text=text)])
